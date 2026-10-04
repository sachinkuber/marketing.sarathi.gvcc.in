"""Builds docs/sachin-to-provide.xlsx from docs/sachin-to-provide.md.

The Markdown file holds the wording of each item. The spreadsheet is where the
owner fills in Status, answer and date. If the spreadsheet already exists, those
three columns are read first and carried over by ID, so rebuilding never loses
what was filled in.

Usage: python3 to_provide_xlsx.py [source.md] [output.xlsx]
"""
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs" / "sachin-to-provide.md"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "docs" / "sachin-to-provide.xlsx"

STATUSES = ["Open", "In progress", "Done", "Not yet due", "Not applicable"]
FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # cells the owner fills in
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def clean(text):
    return re.sub(r"\*\*|`", "", text).strip()


def parse(md):
    """Returns {section title: [rows]} where each row is a list of cell strings."""
    sections, title = {}, None
    for line in md.splitlines():
        m = re.match(r"^## \d+\. (.*)", line)
        if m:
            title = m.group(1).strip()
            sections[title] = []
            continue
        if title and line.startswith("|") and not re.match(r"^\|[-| ]+\|$", line):
            sections[title].append([clean(c) for c in line.strip().strip("|").split("|")])
    return sections


def id_key(item_id):
    m = re.match(r"([A-Z]+)(\d+)", item_id)
    return (m.group(1), int(m.group(2)))


def existing_answers(path):
    """Status, answer and date already filled in, keyed by item ID."""
    if not path.exists():
        return {}
    ws = load_workbook(path)["To provide"]
    kept = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            kept[row[0]] = {"status": row[5], "answer": row[6], "date": row[7]}
    return kept


def style_header(ws, headers, widths):
    for col, (head, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col, value=head)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def body_cell(ws, row, col, value, fill=None, bold=False):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(name=FONT, size=10, bold=bold)
    cell.alignment = WRAP
    cell.border = BORDER
    if fill:
        cell.fill = fill
    return cell


def main():
    sections = parse(SRC.read_text(encoding="utf-8"))
    kept = existing_answers(OUT)
    wb = Workbook()

    # --- To provide -------------------------------------------------------
    ws = wb.active
    ws.title = "To provide"
    headers = ["ID", "Needed when", "What I need", "Exactly what to provide", "Why I need it",
               "Status", "Sachin's answer", "Date answered"]
    style_header(ws, headers, [7, 22, 30, 52, 52, 15, 45, 14])

    items = []
    for title, rows in sections.items():
        if not rows or rows[0][0] != "ID":
            continue
        when = re.sub(r"^Needed ", "", title)
        when = when[0].upper() + when[1:]
        for r in rows[1:]:
            if rows[0][1] == "Item":  # section 8: item, recommended default, why, status
                item_id, what, exact, why, status = r[0], r[1], "Confirm, or say what to change. Recommended default: " + r[2], r[3], r[4]
            else:
                item_id, what, exact, why, status = r[0], r[1], r[2], r[3], r[4]
            items.append((item_id, when, what, exact, why, status))
    items.sort(key=lambda it: id_key(it[0]))

    for n, (item_id, when, what, exact, why, status) in enumerate(items, start=2):
        prior = kept.get(item_id, {})
        body_cell(ws, n, 1, item_id, bold=True)
        body_cell(ws, n, 2, when)
        body_cell(ws, n, 3, what)
        body_cell(ws, n, 4, exact)
        body_cell(ws, n, 5, why)
        body_cell(ws, n, 6, prior.get("status") or status, fill=INPUT_FILL)
        body_cell(ws, n, 7, prior.get("answer"), fill=INPUT_FILL)
        date_cell = body_cell(ws, n, 8, prior.get("date"), fill=INPUT_FILL)
        date_cell.number_format = "DD-MMM-YYYY"
    last = len(items) + 1

    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
    dv.error, dv.errorTitle = "Pick one of the listed statuses.", "Status"
    ws.add_data_validation(dv)
    dv.add(f"F2:F{last}")

    colours = {"Open": "F8CBAD", "In progress": "FFE699", "Done": "C6E0B4",
               "Not yet due": "D9D9D9", "Not applicable": "EDEDED"}
    for status, colour in colours.items():
        ws.conditional_formatting.add(
            f"F2:F{last}",
            CellIsRule(operator="equal", formula=[f'"{status}"'], fill=PatternFill("solid", bgColor=colour)))
    ws.auto_filter.ref = f"A1:H{last}"

    # --- Read me ----------------------------------------------------------
    rm = wb.create_sheet("Read me", 0)
    rm.column_dimensions["A"].width = 24
    rm.column_dimensions["B"].width = 95
    lines = [
        ("Sachin to Provide", None),
        ("What this is", "Everything the project needs from you: access, accounts, facts, decisions and confirmations."),
        ("Where to type", "Sheet 'To provide', the three shaded columns only: Status, Sachin's answer, Date answered. Leave the other columns as they are; they are rebuilt from the document."),
        ("Status", "Pick from the list in the cell: " + ", ".join(STATUSES) + "."),
        ("Filtering", "Use the arrows in the header row to show, for example, only Open items or only items needed now."),
        ("Secrets", "Never type a password, API key or token into this file. For those items write 'handed over' and see the 'Secrets' sheet for how."),
        ("Example", "ID A2 | Status: Done | Sachin's answer: A record added, points to 72.62.195.236 | Date answered: 05-Oct-2026"),
        ("Rebuilding", "This file can be rebuilt from docs/sachin-to-provide.md. What you typed in the shaded columns is kept."),
    ]
    for r, (label, text) in enumerate(lines, start=1):
        a = rm.cell(row=r, column=1, value=label)
        a.font = Font(name=FONT, bold=True, size=14 if r == 1 else 10)
        a.alignment = WRAP
        if text:
            b = rm.cell(row=r, column=2, value=text)
            b.font = Font(name=FONT, size=10)
            b.alignment = WRAP

    start = len(lines) + 2
    rm.cell(row=start, column=1, value="Progress").font = Font(name=FONT, bold=True, size=12)
    for k, status in enumerate(STATUSES, start=start + 1):
        rm.cell(row=k, column=1, value=status).font = Font(name=FONT, size=10)
        c = rm.cell(row=k, column=2, value=f"=COUNTIF('To provide'!$F$2:$F${last},A{k})")
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(horizontal="left")
    total_row = start + 1 + len(STATUSES)
    rm.cell(row=total_row, column=1, value="Total items").font = Font(name=FONT, bold=True, size=10)
    t = rm.cell(row=total_row, column=2, value=f"=COUNTA('To provide'!$A$2:$A${last})")
    t.font = Font(name=FONT, bold=True, size=10)
    t.alignment = Alignment(horizontal="left")

    # --- Secrets and Already provided ---------------------------------------
    for sheet_name, title, widths in [
        ("Secrets", "How to hand over secrets", [34, 100]),
        ("Already provided", "Already provided", [30, 70, 16]),
    ]:
        rows = sections.get(title, [])
        if not rows:
            continue
        sh = wb.create_sheet(sheet_name)
        style_header(sh, rows[0], widths)
        for n, r in enumerate(rows[1:], start=2):
            for col, value in enumerate(r, start=1):
                body_cell(sh, n, col, value)

    wb.save(OUT)
    print(f"wrote {OUT} ({len(items)} items, {len(kept)} answers carried over)")


if __name__ == "__main__":
    main()
